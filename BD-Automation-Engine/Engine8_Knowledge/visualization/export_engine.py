"""
Phase 27A — Export Engine

Export org charts to multiple formats: SVG, PNG, PDF, DOCX, JSON, Mermaid.
"""

import re
from typing import Any, Dict

import structlog

logger = structlog.get_logger(__name__)


class ExportEngine:
    """Export org charts to multiple formats."""

    async def to_svg(self, html: str) -> bytes:
        """Extract SVG from rendered D3.js HTML."""
        # Extract SVG content from HTML
        svg_match = re.search(r"(<svg[^>]*>.*?</svg>)", html, re.DOTALL)
        if svg_match:
            svg_content = svg_match.group(1)
            # Add XML header
            svg_full = '<?xml version="1.0" encoding="UTF-8"?>\n' + svg_content
            return svg_full.encode("utf-8")

        # Fallback: return basic SVG
        return b'<?xml version="1.0"?><svg xmlns="http://www.w3.org/2000/svg"><text x="10" y="20">No SVG content</text></svg>'

    async def to_png(self, html: str, width: int = 1200, height: int = 800) -> bytes:
        """Render HTML to PNG using Playwright (if available)."""
        try:
            from playwright.async_api import async_playwright

            async with async_playwright() as pw:
                browser = await pw.chromium.launch(headless=True)
                page = await browser.new_page(
                    viewport={"width": width, "height": height}
                )
                await page.set_content(html)
                await page.wait_for_timeout(2000)  # Wait for D3.js rendering
                screenshot = await page.screenshot(type="png", full_page=True)
                await browser.close()
                return screenshot
        except ImportError:
            logger.warning("playwright_not_installed")
            return b""
        except Exception as exc:
            logger.error("png_export_error", error=str(exc))
            return b""

    async def to_pdf(self, html: str) -> bytes:
        """Render to PDF with proper page sizing."""
        try:
            from playwright.async_api import async_playwright

            async with async_playwright() as pw:
                browser = await pw.chromium.launch(headless=True)
                page = await browser.new_page()
                await page.set_content(html)
                await page.wait_for_timeout(2000)
                pdf_bytes = await page.pdf(
                    format="A4",
                    landscape=True,
                    print_background=True,
                )
                await browser.close()
                return pdf_bytes
        except ImportError:
            logger.warning("playwright_not_installed")
            return b""
        except Exception as exc:
            logger.error("pdf_export_error", error=str(exc))
            return b""

    async def to_docx(self, org_chart) -> bytes:
        """Generate Word doc with contact table."""
        try:
            from docx import Document

            doc = Document()
            doc.add_heading(org_chart.title, level=1)
            doc.add_paragraph(f"Generated: {org_chart.generated_at}")
            doc.add_paragraph(f"Mode: {org_chart.mode} | Nodes: {len(org_chart.nodes)}")

            # Add table
            if org_chart.nodes:
                table = doc.add_table(rows=1, cols=5)
                table.style = "Table Grid"
                headers = ["Name", "Title", "Company", "Tier", "Reports To"]
                for i, h in enumerate(headers):
                    table.rows[0].cells[i].text = h

                for person in org_chart.nodes:
                    row = table.add_row()
                    row.cells[0].text = person.name
                    row.cells[1].text = person.title
                    row.cells[2].text = person.company
                    row.cells[3].text = str(person.tier)
                    row.cells[4].text = person.reports_to or ""

            import io

            buffer = io.BytesIO()
            doc.save(buffer)
            return buffer.getvalue()
        except ImportError:
            logger.warning("python-docx_not_installed")
            return b""
        except Exception as exc:
            logger.error("docx_export_error", error=str(exc))
            return b""

    async def to_json(self, org_chart) -> Dict[str, Any]:
        """Raw org chart data as JSON."""
        from dataclasses import asdict

        return asdict(org_chart)

    async def to_mermaid(self, org_chart) -> str:
        """Mermaid graph syntax for embedding in markdown."""
        lines = ["graph TD"]

        for person in org_chart.nodes:
            safe_id = person.name.replace(" ", "_").replace("-", "_")
            label = f"{person.name}<br/>{person.title}"
            lines.append(f'    {safe_id}["{label}"]')

        for edge in org_chart.edges:
            source = edge["source"].replace(" ", "_").replace("-", "_")
            target = edge["target"].replace(" ", "_").replace("-", "_")
            lines.append(f"    {source} --> {target}")

        return "\n".join(lines)
