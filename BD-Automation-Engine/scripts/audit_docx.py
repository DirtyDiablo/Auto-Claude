"""Audit the DOCX structure and formatting."""
from docx import Document
from pathlib import Path

path = Path(r"C:\Auto-Claud\Auto-Claude\BD-Automation-Engine\data\deliverables\reports\PTS_BD_EXECUTION_STRATEGY_2026.docx")
doc = Document(str(path))

print("=== DOCUMENT STRUCTURE AUDIT ===")
print()

headings = {1:0, 2:0, 3:0, 4:0}
paragraphs = 0
bullets = 0
numbered = 0

for p in doc.paragraphs:
    name = p.style.name
    if name == "Heading 1": headings[1] += 1
    elif name == "Heading 2": headings[2] += 1
    elif name == "Heading 3": headings[3] += 1
    elif name == "Heading 4": headings[4] += 1
    elif "List Bullet" in name: bullets += 1
    elif "List Number" in name: numbered += 1
    else: paragraphs += 1

tables = len(doc.tables)

print(f"Heading 1 (Phase titles):  {headings[1]}")
print(f"Heading 2 (Sections):      {headings[2]}")
print(f"Heading 3 (Subsections):   {headings[3]}")
print(f"Heading 4 (Sub-sub):       {headings[4]}")
print(f"Body paragraphs:           {paragraphs}")
print(f"Bullet points:             {bullets}")
print(f"Numbered items:            {numbered}")
print(f"Tables:                    {tables}")
print()

print("=== TABLE DETAILS ===")
for i, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    header = [cell.text[:25] for cell in table.rows[0].cells]
    print(f"Table {i+1}: {rows}rows x {cols}cols | {header}")
print()

print("=== HEADING STRUCTURE ===")
for p in doc.paragraphs:
    if p.style.name.startswith("Heading"):
        level = p.style.name.replace("Heading ", "")
        indent = "  " * (int(level) - 1)
        print(f"{indent}H{level}: {p.text[:70]}")
print()

print("=== STYLE VERIFICATION ===")
for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
    s = doc.styles[style_name]
    f = s.font
    color = f.color.rgb if f.color and f.color.rgb else "inherited"
    print(f"{style_name}: font={f.name}, size={f.size}, bold={f.bold}, color={color}")

print()
print("=== TITLE PAGE (first content) ===")
first_texts = [p.text for p in doc.paragraphs[:15] if p.text.strip()]
for t in first_texts:
    print(f'  "{t[:80]}"')

print()
print("=== POTENTIAL ISSUES ===")
issues = []
empty_tables = sum(1 for t in doc.tables if len(t.rows) <= 1)
if empty_tables:
    issues.append(f"{empty_tables} table(s) with only header row (no data)")

for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("| ") and "|" in txt[2:]:
        issues.append(f'Unparsed table row: "{txt[:60]}"')
    if txt.startswith("# ") and p.style.name == "Normal":
        issues.append(f'Raw markdown heading: "{txt[:60]}"')

if not issues:
    print("  None detected - document looks clean!")
else:
    for issue in issues:
        print(f"  WARNING: {issue}")
