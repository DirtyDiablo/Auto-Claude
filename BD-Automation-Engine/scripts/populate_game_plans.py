"""
Populate Game Plan workbook sheets 1-5 with George Maranville's 5 assigned programs.
Uses openpyxl with keep_vba=True to preserve macros.
"""

import openpyxl
from pathlib import Path

WORKBOOK_PATH = Path(__file__).parent.parent / "Game Plan_Blank_GOV.xlsm"

# ── Program Data ──────────────────────────────────────────────────────────────

PROGRAMS = {
    "1": {
        "header": {
            "E6": "ISEE",
            "E7": "",  # City - TBD
            "E8": "DIA",
            "E9": "E-SITE IDIQ",
            "E10": "Oct 2026 (bridge through Aug 29, 2026)",
            "E11": "$100M",
            "E12": 70,
            "E13": "",  # PTS contractors - TBD
            "N6": "2/18/2026",
            "N7": "Q1",
            "N8": 70,
            "N9": "Active",
            "N10": "",  # SB Requirement - TBD
            "N11": "",  # Current Spread - TBD
        },
        "questions": {
            "D16": "DIA infrastructure support. VDI Engineers, Exchange Engineers (L5), Active Directory Engineers, Network Lab Engineers. Team of ~70 contractors supporting DIA IT operations.",
            "D18": "E-SITE IDIQ, direct subcontract to GDIT. Year 5, bridge contract through August 29, 2026.",
            "D20": "TBD - CIS recompete pending. Need to identify competitors on recompete.",
            "D22": "Eileen Esainko (Program Manager), Robyn Moses (Engineering Manager), Jennifer Turnbull (Engineering Manager), Michael Mcknight (Director/Capture Manager). Prime: GDIT.",
            "D24": "Bridge contract currently active through Aug 2026. Recompete expected under CIS vehicle.",
        },
        "contacts": [
            {"name": "Eileen Esainko", "title": "Program Manager", "dept": "ISEE/DIA", "type": "Prime PM", "objectives": "Quarterly touchpoint, staffing needs, recompete timeline"},
            {"name": "Robyn Moses", "title": "Engineering Manager", "dept": "ISEE/DIA", "type": "Prime Eng Mgr", "objectives": "Technical requirements, skill gap identification"},
            {"name": "Jennifer Turnbull", "title": "Engineering Manager", "dept": "ISEE/DIA", "type": "Prime Eng Mgr", "objectives": "Team expansion needs, upcoming openings"},
            {"name": "Michael Mcknight", "title": "Director/Capture Manager", "dept": "GDIT Capture", "type": "Capture Lead", "objectives": "Recompete strategy, teaming opportunities"},
            {"name": "John Allender", "title": "Program Manager", "dept": "ISEE/DIA", "type": "Prime PM", "objectives": "Relationship building, staffing pipeline"},
            {"name": "Jeffrey Watson", "title": "Program Manager", "dept": "ISEE/DIA", "type": "Prime PM", "objectives": "Requirements and upcoming tasking"},
            {"name": "Marisa Rivas", "title": "FSO", "dept": "GDIT Security", "type": "Security", "objectives": "Clearance processing, onboarding support"},
            {"name": "Gabby Ventura", "title": "Subcontracts Manager", "dept": "GDIT Subcontracts", "type": "Subcontracts", "objectives": "Contract terms, new PO requests, rate negotiations"},
        ],
    },
    "2": {
        "header": {
            "E6": "NARSIL",
            "E7": "Bethesda / McLean / Reston",
            "E8": "ODNI / IC",
            "E9": "80GSFC19C0063",
            "E10": "",  # Classified
            "E11": "Classified (~$1.4-1.6B)",
            "E12": "",  # Total contractors - classified
            "E13": "",  # PTS contractors - TBD
            "N6": "2/18/2026",
            "N7": "Q1",
            "N8": "",  # Classified
            "N9": "Active - Expanding",
            "N10": "",  # SB Requirement - TBD
            "N11": "",  # Current Spread - TBD
        },
        "questions": {
            "D16": "13-14 offices, 10 Tasks. TASK2 (Comms Support, 42-43 people: 18 contractors + 24 Peraton FTEs), TASK5, TASK6, TASK8 (largest task). Roles: ML/AI Lead, DevOps Lead, Data Scientists, Task Managers, Software Engineers.",
            "D18": "Direct subcontract to Peraton. Prime contract 80GSFC19C0063.",
            "D20": "Booz Allen Hamilton (disqualified on prior bid for Small Business subcontracting issues). Limited competitive landscape.",
            "D22": "Rondell Shields (Task 2 Lead, 703-275-3636), Arturo Flores (Task 8 Sr PM), Shawn Dimitriades (Director). Prime: Peraton.",
            "D24": "CRITICAL: ODNI will NOT accept previous ODNI employees. Locations: Bethesda (ICCP), Liberty Crossing/McLean, Reston (moving to McLean by end of year). Sha'Rese Davis will not work with PTS.",
        },
        "contacts": [
            {"name": "Rondell Shields", "title": "Task 2 Lead", "dept": "NARSIL/ODNI", "type": "Task Lead", "objectives": "Staffing needs for Task 2, contractor pipeline", "notes": "703-275-3636"},
            {"name": "Arturo Flores", "title": "Task 8 Sr Program Manager", "dept": "NARSIL/ODNI", "type": "Sr PM", "objectives": "Task 8 expansion, headcount growth areas"},
            {"name": "Shawn Dimitriades", "title": "Director", "dept": "Peraton NARSIL", "type": "Director", "objectives": "Overall program strategy, teaming relationship"},
            {"name": "Bob Buckman", "title": "Partnership Manager", "dept": "Peraton", "type": "Partnership", "objectives": "Partnership terms, new task opportunities"},
            {"name": "Victoria Sukar", "title": "", "dept": "NARSIL", "type": "Contact", "objectives": "Relationship building"},
            {"name": "Sha'Rese Davis", "title": "", "dept": "NARSIL", "type": "Contact", "objectives": "N/A - will not work with PTS", "notes": "Will not work with PTS"},
            {"name": "Janelle Barnard", "title": "", "dept": "NARSIL", "type": "Contact", "objectives": "Relationship building, staffing coordination"},
        ],
    },
    "3": {
        "header": {
            "E6": "IRONHIDE",
            "E7": "",  # TBD
            "E8": "DoD / IC",
            "E9": "",  # Verify vehicle
            "E10": "",  # TBD
            "E11": "Unknown",
            "E12": "",  # TBD
            "E13": "",  # TBD
            "N6": "2/18/2026",
            "N7": "Q1",
            "N8": "",  # TBD
            "N9": "Active",
            "N10": "",  # TBD
            "N11": "",  # TBD
        },
        "questions": {
            "D16": "Limited intel - verify departments, skill sets, and headcount. DoD/IC infrastructure program.",
            "D18": "Direct subcontract to GDIT (verify specific contract vehicle).",
            "D20": "",  # TBD
            "D22": "Jim Cemelli (VP), Marvin Lugo (VP - also over Ratchet program), Mike Faulkner (Sr PM), Christopher McGrath (Deputy PM), Jeffrey McClelland (Ops Mgr), Jason Collinsworth (Deputy PM), Andrew Witowski (Systems Engineering Mgr). Prime: GDIT.",
            "D24": "Marvin Lugo also oversees the Ratchet program. Need to develop deeper intel on this contract.",
        },
        "contacts": [
            {"name": "Jim Cemelli", "title": "Vice President", "dept": "GDIT", "type": "Executive", "objectives": "Executive relationship, program overview, growth areas"},
            {"name": "Marvin Lugo", "title": "Vice President", "dept": "GDIT", "type": "Executive", "objectives": "Program strategy, also oversees Ratchet program", "notes": "Also over Ratchet"},
            {"name": "Mike Faulkner", "title": "Senior Program Manager", "dept": "IRONHIDE", "type": "Sr PM", "objectives": "Staffing needs, upcoming requirements"},
            {"name": "Christopher McGrath", "title": "Deputy Program Manager", "dept": "IRONHIDE", "type": "Deputy PM", "objectives": "Day-to-day operations, staffing pipeline"},
            {"name": "Jeffrey McClelland", "title": "Operations Manager", "dept": "IRONHIDE", "type": "Ops Mgr", "objectives": "Operational requirements, contractor onboarding"},
            {"name": "Jason Collinsworth", "title": "Deputy Program Manager", "dept": "IRONHIDE", "type": "Deputy PM", "objectives": "Technical requirements, team structure"},
            {"name": "Andrew Witowski", "title": "Systems Engineering Manager", "dept": "IRONHIDE", "type": "Eng Mgr", "objectives": "Technical skill requirements, engineering openings"},
        ],
    },
    "4": {
        "header": {
            "E6": "TITAN X",
            "E7": "",  # TBD
            "E8": "Space Force",
            "E9": "",  # TBD
            "E10": "",  # TBD
            "E11": "Unknown",
            "E12": "",  # TBD
            "E13": "",  # TBD
            "N6": "2/18/2026",
            "N7": "Q1",
            "N8": "",  # TBD
            "N9": "Established",
            "N10": "",  # TBD
            "N11": "",  # TBD
        },
        "questions": {
            "D16": "Minimal data available. Space domain awareness / satellite program under Space Force. Need to identify departments, skill sets, and headcount.",
            "D18": "Prime: AFS (verify contract vehicle and subcontract arrangement).",
            "D20": "",  # TBD
            "D22": "No key contacts identified yet. Need to develop organizational chart for AFS/Space Force program.",
            "D24": "No Bullhorn data currently available. New intelligence development required.",
        },
        "contacts": [],
    },
    "5": {
        "header": {
            "E6": "RATTLER",
            "E7": "",  # TBD
            "E8": "DoD / IC",
            "E9": "",  # TBD
            "E10": "",  # TBD
            "E11": "Unknown",
            "E12": "",  # TBD
            "E13": "",  # TBD
            "N6": "2/18/2026",
            "N7": "Q1",
            "N8": "",  # TBD
            "N9": "New Add",
            "N10": "",  # TBD
            "N11": "",  # TBD
        },
        "questions": {
            "D16": "New program - all details need to be verified. Not in manual roster.",
            "D18": "Direct subcontract to GDIT (verify specific contract vehicle).",
            "D20": "",  # TBD
            "D22": "No confirmed contacts yet. Verify organizational structure with GDIT.",
            "D24": "Newly added program. All intelligence to be developed. Prime: GDIT.",
        },
        "contacts": [],
    },
}


def populate_sheet(ws, data):
    """Populate a single game plan sheet with program data."""
    # ── Header fields ──
    for cell_ref, value in data["header"].items():
        ws[cell_ref] = value

    # ── Account Review Answers ──
    for cell_ref, value in data["questions"].items():
        ws[cell_ref] = value

    # ── Activity Goals / Contacts (starting at row 29) ──
    for i, contact in enumerate(data["contacts"]):
        row = 29 + i
        ws[f"D{row}"] = contact.get("name", "")
        ws[f"E{row}"] = contact.get("title", "")
        ws[f"F{row}"] = contact.get("dept", "")
        ws[f"H{row}"] = contact.get("type", "")
        ws[f"I{row}"] = contact.get("objectives", "")
        ws[f"N{row}"] = contact.get("notes", "")


def main():
    print(f"Loading workbook: {WORKBOOK_PATH}")
    wb = openpyxl.load_workbook(str(WORKBOOK_PATH), keep_vba=True)

    for sheet_name, data in PROGRAMS.items():
        print(f"Populating sheet '{sheet_name}' with {data['header']['E6']}...")
        ws = wb[sheet_name]
        populate_sheet(ws, data)

    print(f"Saving workbook...")
    wb.save(str(WORKBOOK_PATH))
    print("Done! Workbook saved with macros intact.")

    # Verify by reading back key cells
    print("\n-- Verification --")
    wb2 = openpyxl.load_workbook(str(WORKBOOK_PATH), keep_vba=True)
    for sheet_name in ["1", "2", "3", "4", "5"]:
        ws = wb2[sheet_name]
        prog = ws["E6"].value
        agency = ws["E8"].value
        date = ws["N6"].value
        status = ws["N9"].value
        q1 = ws["D16"].value
        print(f"  Sheet {sheet_name}: {prog} | {agency} | {date} | {status}")
        print(f"    Q1 answer: {str(q1)[:60]}...")


if __name__ == "__main__":
    main()
